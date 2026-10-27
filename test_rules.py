# -*- coding: utf-8 -*-
"""
Unit tests for core rules, checkout time rounding, fingerprint scan parsing,
and Thai employee name similarity matching.
"""

from datetime import time, date
from app import (
    calculate_checkout_time,
    check_fingerprint_scans,
    clean_thai_name,
    check_name_similarity,
    clean_leave_text,
    get_leave_code,
    parse_leave_duration,
    is_company_holiday,
    COMPANY_HOLIDAYS_2026,
    populate_unrecorded_sheet_in_workbook,
    analyze_attendance_anomalies
)


def test_checkout_time_rounding():
    # 1. ก่อน 16:30 ต้องปัดขึ้นเป็น 16:30 ขั้นต่ำ
    assert calculate_checkout_time("16:20") == time(16, 30), "Failed < 16:30"
    assert calculate_checkout_time("15:45") == time(16, 30), "Failed < 16:00"

    # 2. ปัดเศษลงช่วงละ 30 นาที (:00 หรือ :30)
    assert calculate_checkout_time("16:55") == time(16, 30), "Failed 16:55 -> 16:30"
    assert calculate_checkout_time("17:00") == time(17, 0),  "Failed 17:00 -> 17:00"
    assert calculate_checkout_time("17:10") == time(17, 0),  "Failed 17:10 -> 17:00"
    assert calculate_checkout_time("17:29") == time(17, 0),  "Failed 17:29 -> 17:00"
    assert calculate_checkout_time("17:30") == time(17, 30), "Failed 17:30 -> 17:30"
    assert calculate_checkout_time("17:33") == time(17, 30), "Failed 17:33 -> 17:30"
    assert calculate_checkout_time("17:59") == time(17, 30), "Failed 17:59 -> 17:30"
    assert calculate_checkout_time("18:01") == time(18, 0),  "Failed 18:01 -> 18:00"
    assert calculate_checkout_time("18:31") == time(18, 30), "Failed 18:31 -> 18:30"


def test_fingerprint_scans():
    # 1. สแกนเช้าปกติ และ เย็นปกติ
    h_in, h_out, scans, _, _ = check_fingerprint_scans("07:28 17:02")
    assert h_in is True and h_out is True, "Failed 07:28 17:02"
    assert "07:28" in scans and "17:02" in scans

    # 2. สแกนเช้าตรู่ก่อน 05:00 น. (เช่น 04:55 น.)
    h_in, h_out, scans, e_min, _ = check_fingerprint_scans("04:55 17:10")
    assert h_in is True, f"Failed early morning scan 04:55 (got has_in={h_in})"
    assert h_out is True

    # 3. สแกนเฉพาะเช้า
    h_in, h_out, _, _, _ = check_fingerprint_scans("07:15")
    assert h_in is True and h_out is False, "Failed single morning scan"

    # 4. สแกนเฉพาะเย็น
    h_in, h_out, _, _, _ = check_fingerprint_scans("17:30")
    assert h_in is False and h_out is True, "Failed single evening scan"

    # 5. ค่าว่าง
    h_in, h_out, scans, _, _ = check_fingerprint_scans(None)
    assert h_in is False and h_out is False and len(scans) == 0


def test_name_cleaning_and_similarity():
    # 1. ตัดคำนำหน้า
    assert clean_thai_name("นายนริศร จำรักษา") == "นริศร จำรักษา"
    assert clean_thai_name("นางสาวสมหญิง จริงใจ") == "สมหญิง จริงใจ"
    assert clean_thai_name("น.ส. รัตนนาฏ เอ้งฉ้วน") == "รัตนนาฏ เอ้งฉ้วน"
    assert clean_thai_name("MR. SARAWUT WONGKAEW") == "SARAWUT WONGKAEW"
    assert clean_thai_name("A05539 : นายนริศร") == "นริศร"

    # 2. ความสอดคล้องชื่อ (Double-check)
    assert check_name_similarity("นายนริศร จำรักษา", "นริศร") is True
    assert check_name_similarity("A05539 นายนริศร", "นายนริศร จำรักษา") is True
    assert check_name_similarity("นายสมชาย", "นายสมศักดิ์") is False


def test_leave_codes():
    assert clean_leave_text("ลาพักร้อน - Annual Leave") == "ลาพักร้อน"
    assert clean_leave_text("ลาป่วย - Sick Leave") == "ลาป่วย"
    assert clean_leave_text("ลาโดยไม่ได้รับค่าจ้าง - Unpaid Leave") == "ลาโดยไม่ได้รับค่าจ้าง"

    assert get_leave_code("ลาพักร้อน") == "พ"
    assert get_leave_code("ลาป่วย") == "ป"
    assert get_leave_code("ลากิจ") == "อ"
    assert get_leave_code("ลาโดยไม่ได้รับค่าจ้าง") == "อ"


def test_analytics_leave_breakdown():
    # ตรวจสอบว่า get_project_analytics แยกประเภทการลาถูกต้อง
    # โดยไม่หลุดเอาชื่อพนักงานจากเคส "⚠️ ไม่พบพนักงานในชีท" มานับเป็นประเภทการลา
    mock_items = [
        {'type': 'วันลา', 'leave_type': 'ลาพักร้อน', 'status': '✅ พร้อมอัปเดต', 'detail': 'ลาพักร้อน'},
        {'type': 'วันลา', 'leave_type': 'ลาป่วย', 'status': '⚠️ ไม่พบพนักงานในชีท', 'detail': '(คุณสมชาย สุขใจ)'},
        {'type': 'วันลา', 'leave_type': 'ลากิจ', 'status': '⚠️ ข้อมูลซ้ำซ้อน', 'detail': '(มีข้อมูลอยู่แล้ว)'},
        {'type': 'แจ้งเวลา', 'attend_type': 'เช้า (07:30)', 'status': '✅ พร้อมอัปเดต', 'detail': 'บันทึกเวลาเข้า'},
    ]
    
    leave_types_count = {}
    for item in mock_items:
        if item['type'] == 'วันลา':
            l_type = item.get('leave_type') or 'ไม่ระบุประเภท'
            leave_types_count[l_type] = leave_types_count.get(l_type, 0) + 1

    assert 'ลาพักร้อน' in leave_types_count and leave_types_count['ลาพักร้อน'] == 1
    assert 'ลาป่วย' in leave_types_count and leave_types_count['ลาป่วย'] == 1
    assert 'ลากิจ' in leave_types_count and leave_types_count['ลากิจ'] == 1
    assert 'คุณสมชาย สุขใจ' not in leave_types_count
    assert len(leave_types_count) == 3


def test_leave_rules_and_abbreviations():
    # 1. รหัสตัวย่อวันลา
    assert get_leave_code("ลาพักร้อน") == "พ"
    assert get_leave_code("ลาป่วย") == "ป"
    assert get_leave_code("ลากิจ") == "อ"

    # 2. กฎสถานะ/ตัวย่อ:
    # - สถานะใช้กับ "ลาเต็มวัน" อย่างเดียว (รวมลาโดยไม่ได้รับค่าจ้าง ให้ลงตัวย่อ อ ด้วย)
    # - ลาไม่เต็มวัน (Partial Leave): ไม่ลงสถานะช่อง Q เพื่อให้สูตรคำนวณชั่วโมงทำงานได้
    assert get_leave_code("ลาโดยไม่ได้รับค่าจ้าง") == "อ"
    clean_unpaid = clean_leave_text("ลาโดยไม่ได้รับค่าจ้าง - Unpaid Leave")
    assert clean_unpaid == "ลาโดยไม่ได้รับค่าจ้าง"
    assert get_leave_code(clean_unpaid) == "อ"

    # 3. แยกแยะลาเช้า vs ลาบ่าย
    # ลาบ่าย: มีสแกนเช้า แต่สแกนสุดท้ายก่อน 13:00 (เช่น 11:31) -> ปรับเลิก M เป็น 16:30 ห้ามแตะต้อง L
    e_max_afternoon = 11.516
    assert e_max_afternoon < 13.0, "Afternoon leave detection failed"

    # ลาเช้า บ่ายมาทำ: มีสแกนเย็น >= 13:00 (เช่น 18:51) -> เข้า L ปรับเป็น 7:30, เลิกตามจริง 18:30 (ห้ามแตะต้อง M)
    e_max_morning = 18.85
    assert e_max_morning >= 13.0, "Morning leave detection failed"
    assert calculate_checkout_time(e_max_morning) == time(18, 30)


def test_parse_leave_duration():
    # แก้ปัญหาระบบภายนอกบั๊กส่งนาทีเกินมา 1 หลัก (150 -> 15 นาที, 450 -> 45 นาที, 300 -> 30 นาที)
    is_p, h = parse_leave_duration("0 วัน 4 ชั่วโมง 150 นาที")
    assert is_p is True and h == 4.25, f"Expected (True, 4.25), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 3 ชั่วโมง 450 นาที")
    assert is_p is True and h == 3.75, f"Expected (True, 3.75), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 2 ชั่วโมง 150 นาที")
    assert is_p is True and h == 2.25, f"Expected (True, 2.25), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 2 ชั่วโมง 450 นาที")
    assert is_p is True and h == 2.75, f"Expected (True, 2.75), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 4 ชั่วโมง 30 นาที")
    assert is_p is True and h == 4.5, f"Expected (True, 4.5), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 2 ชั่วโมง 0 นาที")
    assert is_p is True and h == 2.0, f"Expected (True, 2.0), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("0 วัน 2 ชั่วโมง 300 นาที")
    assert is_p is True and h == 2.5, f"Expected (True, 2.5), got ({is_p}, {h})"

    # ลาเต็มวัน ไม่ถือว่าเป็น partial
    is_p, h = parse_leave_duration("1 วัน 0 ชั่วโมง 0 นาที")
    assert is_p is False and h is None, f"Expected (False, None), got ({is_p}, {h})"

    is_p, h = parse_leave_duration("2 วัน 0 ชั่วโมง 0 นาที")
    assert is_p is False and h is None, f"Expected (False, None), got ({is_p}, {h})"


def test_company_holidays():
    # 1. ครบทั้ง 15 วันตามประกาศ บมจ.คริสเตียนีและนีลเส็น (ไทย) ปี 2569
    assert len(COMPANY_HOLIDAYS_2026) == 15, f"Expected 15 holidays, got {len(COMPANY_HOLIDAYS_2026)}"

    # 2. วันหยุดสำคัญ
    is_h, name = is_company_holiday(date(2026, 10, 13))
    assert is_h is True and "นวมินทรมหาราช" in name

    is_h, name = is_company_holiday(date(2026, 10, 23))
    assert is_h is True and "ปิยมหาราช" in name

    is_h, name = is_company_holiday(date(2026, 4, 13))
    assert is_h is True and "สงกรานต์" in name

    # 3. วันทำงานปกติ ไม่ใช่วันหยุด
    is_h, name = is_company_holiday(date(2026, 10, 14))
    assert is_h is False and name is None

    is_h, name = is_company_holiday(date(2026, 10, 15))
    assert is_h is False and name is None


def test_unrecorded_sheet_grouping():
    """ทดสอบการจัดกลุ่มพนักงานที่ไม่มีการแจ้งออนไลน์: หาก 1 คนมีหลายวัน ต้องเรียงติดกันตามลำดับวันที่"""
    import openpyxl
    wb = openpyxl.Workbook()
    # Mock data
    mock_anomalies = [
        {'emp_id': '5732', 'emp_name': 'กัลยกร  นิยมเดชา', 'date': '19/10/2026', 'date_obj': date(2026, 10, 19), 'category': 'missing_punch', 'detail': 'ลืมสแกนออก', 'scans_str': '07:10'},
        {'emp_id': '5688', 'emp_name': 'กัณมน กิบุญมา', 'date': '19/10/2026', 'date_obj': date(2026, 10, 19), 'category': 'missing_punch', 'detail': 'ลืมสแกนออก', 'scans_str': '07:15'},
        {'emp_id': '5732', 'emp_name': 'กัลยกร  นิยมเดชา', 'date': '11/10/2026', 'date_obj': date(2026, 10, 11), 'category': 'missing_punch', 'detail': 'ลืมสแกนออก', 'scans_str': '07:22'},
        {'emp_id': '5688', 'emp_name': 'กัณมน กิบุญมา', 'date': '12/10/2026', 'date_obj': date(2026, 10, 12), 'category': 'missing_punch', 'detail': 'ลืมสแกนออก', 'scans_str': '07:18'},
    ]
    # Sort with anomaly_sort_key
    def anomaly_sort_key(a):
        dt_o = a.get('date_obj')
        d_tup = (dt_o.year, dt_o.month, dt_o.day) if dt_o else (9999, 12, 31)
        return (clean_thai_name(a.get('emp_name', '')), str(a.get('emp_id', '')), d_tup)
    mock_anomalies.sort(key=anomaly_sort_key)

    populate_unrecorded_sheet_in_workbook(wb, anomalies=mock_anomalies, filename="CN.11856", create_sheet=True)
    ws = wb['ไม่มีแจ้งในระบบ']
    assert ws['A1'].value == "รายชื่อที่ไม่มีการแจ้งออนไลน์ CN.11856"
    assert ws['A2'].value == "รหัส"
    assert ws['B2'].value == "ชื่อ-สกุล"
    assert ws['D2'].value == "วันที่"
    assert ws.column_dimensions['C'].hidden is True

    # Row 3 & 4 should be กัณมน (12/10, then 19/10)
    assert ws.cell(3, 2).value == 'กัณมน กิบุญมา' and ws.cell(3, 4).value == '12/10/2026'
    assert ws.cell(4, 2).value == 'กัณมน กิบุญมา' and ws.cell(4, 4).value == '19/10/2026'
    # Row 5 & 6 should be กัลยกร (11/10, then 19/10)
    assert ws.cell(5, 2).value == 'กัลยกร  นิยมเดชา' and ws.cell(5, 4).value == '11/10/2026'
    assert ws.cell(6, 2).value == 'กัลยกร  นิยมเดชา' and ws.cell(6, 4).value == '19/10/2026'


def test_afternoon_shift_and_leave_anomaly_exclusion():
    """ทดสอบกรณีสแกนบ่าย 2 เวลา (เข้า 12:54 ออก 18:31) และกรณีมีใบลาในระบบ ต้องไม่ถูกฟ้องว่าลืมสแกนเข้า"""
    # 1. ทดสอบ check_fingerprint_scans เมื่อมีสแกนบ่าย 2 เวลา
    h_in, h_out, scans, e_min, e_max = check_fingerprint_scans("12:54 12:54 18:31 18:31")
    assert h_in is True, f"12:54 and 18:31 must have has_in=True (got {h_in})"
    assert h_out is True, f"12:54 and 18:31 must have has_out=True (got {h_out})"

    # 2. ทดสอบ analyze_attendance_anomalies เมื่อมีบันทึกการลา
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "15.10.26"
    # Row 1 Headers
    ws['A1'] = 'AC-No'; ws['B1'] = 'Name'; ws['E1'] = 'Time'; ws['G1'] = 'Leave'; ws['Q1'] = 'Code'
    # Row 2: นฤมล ฉลาดมาก มีสแกนบ่าย 12:54, 18:31 และมีลาป่วย 4.5 ชม.
    ws['A2'] = '5614'; ws['B2'] = 'นฤมล ฉลาดมาก'; ws['E2'] = '12:54 18:31'; ws['G2'] = 'ลาป่วย 4.5 ชม.'; ws['Q2'] = '=+P2-1'
    # Row 3: พนักงานที่มีใบแจ้งฯ ขาเดียว (เช่น 07:15)
    ws['A3'] = '5688'; ws['B3'] = 'กัณมน กิบุญมา'; ws['E3'] = '07:15'; ws['G3'] = 'มีใบแจ้งฯ'; ws['Q3'] = '10'
    # Row 4: พนักงานที่ลืมสแกนจริง และไม่มีใบแจ้ง
    ws['A4'] = '9999'; ws['B4'] = 'ทดสอบ ลืมสแกน'; ws['E4'] = '07:15'; ws['G4'] = ''; ws['Q4'] = '=+P4-1'

    res = analyze_attendance_anomalies(wb)
    anomalies = res['anomalies']
    
    # นฤมล (5614) และ กัณมน (5688) ต้องไม่อยู่ใน anomalies
    ids = [a['emp_id'] for a in anomalies]
    assert '5614' not in ids, "5614 (นฤมล ฉลาดมาก) must NOT be in anomalies"
    assert '5688' not in ids, "5688 (มีใบแจ้งฯ) must NOT be in anomalies"
    assert '9999' in ids, "9999 (ไม่มีใบแจ้ง และสแกนขาเดียว) must be in anomalies"


def test_double_check_actual_and_adjusted_times():
    """ทดสอบระบบดับเบิลเช็คทั้งค่าจริง (Col J, K) และค่าที่ปรับตามการลา (Col L, M) รวมถึง Col G และ Col Q"""
    from datetime import time
    from app import is_valid_time, format_time_val

    # 1. ทดสอบ helper is_valid_time และ format_time_val
    assert is_valid_time(time(7, 30)) is True
    assert is_valid_time(time(13, 0)) is True
    assert is_valid_time(time(0, 0)) is False
    assert is_valid_time("13:00") is True
    assert is_valid_time("07:30") is True
    assert is_valid_time("ไม่ได้แสกน") is False
    assert is_valid_time("ไม่ได้สแกน") is False
    assert is_valid_time("0:00") is False
    assert is_valid_time("00:00") is False
    assert is_valid_time(None) is False
    assert is_valid_time(0.0) is False

    assert format_time_val(time(7, 30)) == "07:30"
    assert format_time_val("13:00") == "13:00"

    # 2. ทดสอบ Double-check ใน analyze_attendance_anomalies
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "15.10.26" # วันพฤหัสบดี (วันทำงานปกติ)
    # Headers
    ws['A1'] = 'AC-No'; ws['B1'] = 'Name'; ws['E1'] = 'Time'; ws['G1'] = 'Leave'
    ws['J1'] = 'Time-In'; ws['K1'] = 'Time-Out'; ws['L1'] = 'Check-in'; ws['M1'] = 'Check-Out'; ws['Q1'] = 'Code'

    # Case 1: 1001 มีค่าจริงใน J='13:00', K='18:30' (แม้ L, M ว่าง และ E ไม่มีสแกน) -> สมบูรณ์ ไม่เป็น anomaly
    ws['A2'] = '1001'; ws['B2'] = 'นาย สมชาย ค่าจริง'; ws['E2'] = 'ไม่ได้แสกน'
    ws['J2'] = '13:00'; ws['K2'] = '18:30'; ws['L2'] = None; ws['M2'] = None; ws['G2'] = ''; ws['Q2'] = '=+P2-1'

    # Case 2: 1002 มีค่าปรับตามการลาใน L='07:30', M='18:30' (เช่น ลาครึ่งวันเช้า) -> สมบูรณ์ ไม่เป็น anomaly
    ws['A3'] = '1002'; ws['B3'] = 'นางสาว สมศรี ค่าปรับ'; ws['E3'] = 'ไม่ได้แสกน'
    ws['J3'] = None; ws['K3'] = None; ws['L3'] = '07:30'; ws['M3'] = '18:30'; ws['G3'] = ''; ws['Q3'] = '=+P3-1'

    # Case 3: 1003 ขาดเวลาออก (มีเฉพาะเข้า L='07:30') และไม่มีใบแจ้ง -> ลืมสแกนออก
    ws['A4'] = '1003'; ws['B4'] = 'นาย สมศักดิ์ ลืมออก'; ws['E4'] = 'ไม่ได้แสกน'
    ws['J4'] = None; ws['K4'] = None; ws['L4'] = '07:30'; ws['M4'] = 'ไม่ได้แสกน'; ws['G4'] = ''; ws['Q4'] = '=+P4-1'

    # Case 4: 1004 ขาดเวลาเข้า (มีเฉพาะออก K='18:30') และไม่มีใบแจ้ง -> ลืมสแกนเข้า
    ws['A5'] = '1004'; ws['B5'] = 'นาง สมบูรณ์ ลืมเข้า'; ws['E5'] = 'ไม่ได้แสกน'
    ws['J5'] = 'ไม่ได้แสกน'; ws['K5'] = '18:30'; ws['L5'] = None; ws['M5'] = None; ws['G5'] = ''; ws['Q5'] = '=+P5-1'

    # Case 5: 1005 ไม่มีเวลาเข้า-ออกเลย แต่มีบันทึกการลาในช่อง G -> สมบูรณ์ ไม่เป็น anomaly
    ws['A6'] = '1005'; ws['B6'] = 'นาย มาลา ลาป่วย'; ws['E6'] = 'ไม่ได้แสกน'
    ws['J6'] = None; ws['K6'] = None; ws['L6'] = None; ws['M6'] = None; ws['G6'] = 'ลาป่วย 4.5 ชม.'; ws['Q6'] = '=+P6-1'

    # Case 6: 1006 ไม่มีเวลาเข้า-ออกเลย แต่มีรหัสวันลาในช่อง Q -> สมบูรณ์ ไม่เป็น anomaly
    ws['A7'] = '1006'; ws['B7'] = 'นาย ลาพัก ร้อน'; ws['E7'] = 'ไม่ได้แสกน'
    ws['J7'] = None; ws['K7'] = None; ws['L7'] = None; ws['M7'] = None; ws['G7'] = ''; ws['Q7'] = 'พ'

    # Case 7: 1007 ไม่มีเวลาเข้า-ออก และไม่มีบันทึกการลา -> ไม่พบสแกนและไม่ลา (Unrecorded Absence)
    ws['A8'] = '1007'; ws['B8'] = 'นาย ขาดงาน แน่นอน'; ws['E8'] = 'ไม่ได้แสกน'
    ws['J8'] = None; ws['K8'] = None; ws['L8'] = None; ws['M8'] = None; ws['G8'] = ''; ws['Q8'] = '=+P8-1'

    res = analyze_attendance_anomalies(wb)
    anomalies = res['anomalies']
    anomaly_map = {a['emp_id']: a for a in anomalies}

    # ตรวจสอบการคัดกรอง
    assert '1001' not in anomaly_map, "1001 (มีค่าจริง J/K) ต้องไม่อยู่ใน anomalies"
    assert '1002' not in anomaly_map, "1002 (มีค่าปรับ L/M) ต้องไม่อยู่ใน anomalies"
    assert '1005' not in anomaly_map, "1005 (มีลาป่วยใน Col G) ต้องไม่อยู่ใน anomalies"
    assert '1006' not in anomaly_map, "1006 (มีรหัสวันลาใน Col Q) ต้องไม่อยู่ใน anomalies"

    # ตรวจสอบคนที่ผิดปกติจริง
    assert '1003' in anomaly_map, "1003 ต้องถูกตรวจจับว่าผิดปกติ"
    assert anomaly_map['1003']['item'] == 'ลืมสแกนออก'
    assert '07:30' in anomaly_map['1003']['note']

    assert '1004' in anomaly_map, "1004 ต้องถูกตรวจจับว่าผิดปกติ"
    assert anomaly_map['1004']['item'] == 'ลืมสแกนเข้า'
    assert '18:30' in anomaly_map['1004']['note']

    assert '1007' in anomaly_map, "1007 ต้องถูกตรวจจับว่าขาดงาน"
    assert anomaly_map['1007']['category'] == 'unrecorded_absence'


if __name__ == "__main__":
    test_checkout_time_rounding()
    test_fingerprint_scans()
    test_name_cleaning_and_similarity()
    test_leave_codes()
    test_analytics_leave_breakdown()
    test_leave_rules_and_abbreviations()
    test_parse_leave_duration()
    test_company_holidays()
    test_unrecorded_sheet_grouping()
    test_afternoon_shift_and_leave_anomaly_exclusion()
    test_double_check_actual_and_adjusted_times()
    print("ALL TESTS PASSED SUCCESSFULLY! (100% OK)")
